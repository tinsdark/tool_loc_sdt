"""Loc so dien thoai tu file Word (.docx), Excel (.xlsx) va .txt -> file .txt

Cach dung:
    python loc_sdt.py a.docx b.xlsx c.txt -o ketqua.txt
    python loc_sdt.py                      (khong tham so -> mo cua so chon file)

Output:
    ketqua.txt       : SDT hop le, da chuan hoa 0xxxxxxxxx, bo trung, moi dong 1 so
    ketqua_loi.txt   : cac gia tri trong giong SDT nhung khong hop le (kem ly do + nguon)
"""
import argparse
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

# Dau so di dong VN: 032-039, 052/055/056/058/059, 070/076-079, 081-089, 090-099
MOBILE = re.compile(r"^0(3[2-9]|5[25689]|7[06-9]|8[1-9]|9\d)\d{7}$")
# So 11 chu so: chi can 0 + dau 3/5/7/8/9
MOBILE11 = re.compile(r"^0[35789]\d{9}$")
SKIP = "skip"  # dau so 01/02 -> bo qua, khong bao loi
# Mot "ung vien": bat dau bang so, cho phep chu/dau noi dinh lien; cac cum cach nhau 1 dau cach
CANDIDATE = re.compile(r"\+?\d[\dA-Za-z.\-()]*(?: \d[\dA-Za-z.\-()]*)*")
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


# ---------- doc du lieu: moi ham tra ve (nguon, doan_text) ----------
def read_txt(path):
    for i, line in enumerate(Path(path).read_text(encoding="utf-8-sig", errors="ignore").splitlines(), 1):
        yield f"dong {i}", line


def read_docx(path):
    """Doc moi doan (ke ca trong bang, header/footer, textbox) truc tiep tu XML."""
    with zipfile.ZipFile(path) as z:
        parts = [n for n in z.namelist() if re.match(r"word/(document|header\d*|footer\d*)\.xml$", n)]
        for part in parts:
            root = ET.fromstring(z.read(part))
            for i, p in enumerate(root.iter(W_NS + "p"), 1):
                text = "".join(t.text or "" for t in p.iter(W_NS + "t"))
                if text.strip():
                    yield f"{part.split('/')[-1]} doan {i}", text


def read_xlsx(path):
    from openpyxl import load_workbook  # import tre de file txt/docx khong can openpyxl
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    continue
                if isinstance(v, float) and v.is_integer():
                    v = int(v)
                yield f"{ws.title}!{c.coordinate}", str(v)


READERS = {".txt": read_txt, ".docx": read_docx, ".xlsx": read_xlsx, ".xlsm": read_xlsx}


# ---------- chuan hoa ----------
def normalize(raw):
    """Tra ve (so_chuan_hoa, None) | (None, ly_do) | (None, SKIP) neu dau so 01/02."""
    d = re.sub(r"\D", "", raw)  # chi giu so: bo chu, dau cach, ngoac, cham...
    if d.startswith("84") and len(d) == 11:
        d = "0" + d[2:]
    elif len(d) == 9 and d[0] in "35789":  # Excel an mat so 0 dau
        d = "0" + d
    if d[:2] in ("01", "02") or (d[:1] in "12" and len(d) in (9, 10)):  # 01/02 (ke ca khi mat so 0)
        return None, SKIP
    if len(d) == 10:
        return (d, None) if MOBILE.match(d) else (None, "đầu số không hợp lệ")
    if len(d) == 11:
        return (d, None) if MOBILE11.match(d) else (None, "đầu số không hợp lệ")
    return None, f"sai độ dài ({len(d)} số)"


def extract(text):
    """Yield (so_hop_le | None, raw, ly_do). Chi bao loi khi ung vien co >= 8 chu so."""
    for m in CANDIDATE.finditer(text):
        raw = m.group()
        if sum(ch.isdigit() for ch in raw) < 8:
            continue
        num, why = normalize(raw)
        if num or why == SKIP:
            yield num, raw, why
            continue
        parts = raw.split(" ")  # vd "0912345678 0987654321" dinh nhau trong 1 o
        if len(parts) > 1:
            res = [normalize(p) for p in parts]
            if all(n or w == SKIP for n, w in res):
                for n, w in res:
                    yield n, raw, w
                continue
        yield None, raw, why


def run(inputs, out):
    seen, valid, rejected = set(), [], []
    total = skipped = 0
    for f in inputs:
        reader = READERS.get(Path(f).suffix.lower())
        if not reader:
            print(f"Bo qua (khong ho tro): {f}", file=sys.stderr)
            continue
        for src, text in reader(f):
            for num, raw, why in extract(text):
                if why == SKIP:
                    skipped += 1
                elif num:
                    total += 1
                    if num not in seen:
                        seen.add(num)
                        valid.append(num)
                else:
                    rejected.append(f"{raw}\t{why}\t{Path(f).name}:{src}")
    out = Path(out)
    out.write_text("\n".join(valid) + ("\n" if valid else ""), encoding="utf-8")
    err = out.with_name(out.stem + "_loi.txt")
    err.write_text("\n".join(rejected) + ("\n" if rejected else ""), encoding="utf-8")
    msg = (f"Tim thay {total} SDT hop le -> {len(valid)} so khac nhau (bo {total - len(valid)} trung)\n"
           f"Bo dau so 01/02: {skipped}\n"
           f"Khong hop le: {len(rejected)} (xem {err.name})\nKet qua: {out}")
    print(msg)
    return msg


def gui():
    import tkinter as tk
    from tkinter import filedialog, messagebox
    root = tk.Tk()
    root.withdraw()
    files = filedialog.askopenfilenames(
        title="Chon file Word / Excel / txt",
        filetypes=[("Word/Excel/txt", "*.docx *.xlsx *.xlsm *.txt")])
    if not files:
        return
    out = filedialog.asksaveasfilename(title="Luu ket qua", defaultextension=".txt",
                                       initialfile="ketqua.txt")
    if out:
        messagebox.showinfo("Xong", run(files, out))


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    ap.add_argument("files", nargs="*", help=".docx / .xlsx / .txt")
    ap.add_argument("-o", "--out", default="ketqua.txt")
    a = ap.parse_args()
    run(a.files, a.out) if a.files else gui()
